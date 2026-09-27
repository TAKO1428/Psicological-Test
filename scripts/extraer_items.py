"""
Extrae los items de niveles_de_complejidad.xlsx a un JSON estructurado
que luego se carga a la base de datos (ver db/seed.sql o el endpoint de importación).

Reglas aplicadas:
- Selector 1 / Selector 2 = par de ejemplo (se muestra antes de la pregunta real, sin pedir respuesta)
- Muestra + Eco1..Eco4 = la pregunta real de igualación a la muestra
- Respuesta correcta: se detecta por el relleno VERDE de la celda Eco. Si ninguna
  celda Eco tiene verde, se usa la columna "Respuesta" (J) comparando texto normalizado.
- fase se toma del nombre de la hoja (Preprueba / Entrenamiento / Postprueba)

Uso:
    python3 extraer_items.py /ruta/al/excel.xlsx salida.json
"""
import sys
import json
import unicodedata
import openpyxl

VERDES = {"FF00FF00", "FF93C47D"}  # los dos tonos de verde usados en el archivo
FASE_POR_HOJA = {
    "Preprueba": "preprueba",
    "Entrenamiento": "entrenamiento",
    "Postprueba": "postprueba",
}


def normalizar(texto):
    if texto is None:
        return None
    t = str(texto).strip().lower()
    t = unicodedata.normalize("NFKD", t)
    t = "".join(c for c in t if not unicodedata.combining(c))
    return t


def color_de(celda):
    try:
        fg = celda.fill.fgColor.rgb
        return fg if isinstance(fg, str) else None
    except Exception:
        return None


def extraer_hoja(ws, fase):
    items = []
    incompletos = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10):
        numero, relacion, sel1, sel2, muestra, e1, e2, e3, e4, respuesta_txt = [
            c.value for c in row
        ]
        ecos = [e1, e2, e3, e4]

        # fila totalmente vacía (sin relación siquiera) -> se ignora
        if relacion is None and all(v is None for v in [muestra, e1, e2, e3, e4]):
            continue

        # fila incompleta (falta muestra o algún eco) -> se reporta, no se carga
        if muestra is None or any(e is None for e in ecos):
            incompletos.append(
                {
                    "fila_excel": row[0].row,
                    "fase": fase,
                    "relacion": relacion,
                    "motivo": "faltan datos (muestra o ecos vacíos)",
                }
            )
            continue

        # detectar cuál eco es verde
        celdas_eco = row[5:9]
        indice_correcto = None
        for idx, celda in enumerate(celdas_eco):
            if color_de(celda) in VERDES:
                indice_correcto = idx
                break

        # fallback: comparar texto de la columna Respuesta contra los ecos
        if indice_correcto is None and respuesta_txt is not None:
            resp_norm = normalizar(respuesta_txt)
            for idx, valor in enumerate(ecos):
                if normalizar(valor) == resp_norm:
                    indice_correcto = idx
                    break

        if indice_correcto is None:
            incompletos.append(
                {
                    "fila_excel": row[0].row,
                    "fase": fase,
                    "relacion": relacion,
                    "motivo": "no se pudo determinar respuesta correcta "
                    f"(verde ausente y J={respuesta_txt!r} no coincide con ningún Eco)",
                }
            )
            continue

        items.append(
            {
                "fila_excel": row[0].row,
                "fase": fase,
                "relacion": (relacion or "").strip(),
                "ejemplo_selector_1": (sel1 or "").strip() if sel1 else None,
                "ejemplo_selector_2": (sel2 or "").strip() if sel2 else None,
                "muestra": str(muestra).strip(),
                "ecos": [str(e).strip() for e in ecos],
                "indice_correcto": indice_correcto,  # 0-based, referido a la lista "ecos"
            }
        )

    return items, incompletos


def main():
    if len(sys.argv) != 3:
        print("Uso: python3 extraer_items.py <excel_entrada> <json_salida>")
        sys.exit(1)

    ruta_excel, ruta_salida = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(ruta_excel, data_only=False)

    todos_items = []
    todos_incompletos = []

    for nombre_hoja, fase in FASE_POR_HOJA.items():
        if nombre_hoja not in wb.sheetnames:
            print(f"Aviso: no se encontró la hoja '{nombre_hoja}', se omite.")
            continue
        ws = wb[nombre_hoja]
        items, incompletos = extraer_hoja(ws, fase)
        todos_items.extend(items)
        todos_incompletos.extend(incompletos)
        print(f"{nombre_hoja}: {len(items)} items válidos, {len(incompletos)} omitidos")

    with open(ruta_salida, "w", encoding="utf-8") as f:
        json.dump(
            {"items": todos_items, "omitidos": todos_incompletos},
            f,
            ensure_ascii=False,
            indent=2,
        )

    print(f"\nTotal items válidos: {len(todos_items)}")
    print(f"Total omitidos (incompletos/ambiguos): {len(todos_incompletos)}")
    print(f"Guardado en: {ruta_salida}")


if __name__ == "__main__":
    main()
